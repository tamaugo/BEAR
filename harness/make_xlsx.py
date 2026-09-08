#!/usr/bin/env python3
"""
Converts the pipeline's pipe-delimited text output into a .xlsx workbook.

Usage:
  python3 make_xlsx.py <input.txt> <output.xlsx>

--- Why this script exists ---

Agent 3 is an LLM. It can only write text. A .xlsx file is a binary ZIP of
XML documents, which no language model can emit reliably. So the split is:
Agent 3 writes a plain pipe-delimited text file, and this script turns that
file into a spreadsheet deterministically.

That split is deliberate and this script is deliberately dumb. It makes NO
judgement calls — it does not clean part names, infer missing prices, guess
at a URL, drop rows it dislikes, or reorder anything. Every transformation
rule in this pipeline lives in Agent 3's prompt, exactly one place. If you
find yourself wanting to add "smart" behaviour here, it belongs in Agent 3
instead.

--- Input format ---

One row per line, exactly 4 pipe-separated fields:

    part info | price | url | image

Real examples:

    MAZDA 6 MK2 2008 SEDAN 2.5 PETROL - Off Side Front Bumper Bracket \
GS1D500T1 | 44.99 | https://www.ebay.co.uk/itm/123456789 | img_2225_OSF.JPEG
    FAILED - Agent 2 - No eBay Listing Found | | | img_2240.JPEG

Failure rows have empty price and url fields but ALWAYS carry an image
field, so a failed part is still traceable back to its photograph.

Leading and trailing whitespace is stripped from every field. Completely
blank lines are skipped (they are spacing, not rows).

A line that does not have exactly 4 fields is NOT silently dropped. It is
written into the sheet with the whole raw line in the part info column and
a MALFORMED marker in front of it, and a warning is printed to stderr. A
bad row must be visible in the output, never lost.

--- Output format ---

Column A = part info    B = price    C = url    D = image

Header row is bold and frozen. Price is written as a real NUMBER wherever
the text parses as one, so the sheet can sum and sort it; anything that
does not parse (an empty failure-row price, a value with a currency symbol
or thousands separator) is left as text, unchanged and visible. URLs that
look like http(s) links are made clickable.

--- Dependencies ---

openpyxl. Everything else is standard library. If openpyxl is missing the
script says so and exits rather than failing halfway through.
"""

import os
import sys

try:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit(
        "This script needs the openpyxl package, which is not installed.\n"
        "Install it with:\n"
        "  pip3 install openpyxl\n"
        "(Nothing else is required — the rest of the script is standard library.)"
    )

HEADERS = ["part info", "price", "url", "image"]
FIELD_COUNT = len(HEADERS)

# part info is long, url is medium, price and image are narrow.
COLUMN_WIDTHS = [70, 10, 46, 22]

MALFORMED_PREFIX = "MALFORMED ROW"


def clean(text: str) -> str:
    """Strip surrounding whitespace, and drop the control characters that
    the XLSX format forbids. Stray control bytes are the one thing that
    would make openpyxl refuse to write the file at all, so they are
    removed rather than allowed to abort a whole run. Nothing else about
    the text is touched."""
    return ILLEGAL_CHARACTERS_RE.sub("", text).strip()


def parse_price(text: str):
    """Return a float when the price text is a plain number, otherwise the
    text unchanged.

    Intentionally strict: only what float() accepts is treated as a number.
    "44.99" becomes a number; "", "£44.99" and "1,299.00" stay as text and
    stay visible in the sheet. Guessing at currency symbols or separators
    would be a judgement call, and those do not belong in this script."""
    try:
        return float(text)
    except ValueError:
        return text


def parse_line(raw: str):
    """Turn one input line into a (fields, malformed) pair.

    A well-formed line gives its 4 stripped fields and malformed=False. A
    line with any other number of fields gives the raw line, marked, in the
    part info column with the rest blank, and malformed=True — so it lands
    in the sheet where a human will see it."""
    fields = [clean(f) for f in raw.split("|")]
    if len(fields) == FIELD_COUNT:
        return fields, False

    marker = (
        f"{MALFORMED_PREFIX} (expected {FIELD_COUNT} fields, "
        f"got {len(fields)}): {clean(raw)}"
    )
    return [marker, "", "", ""], True


def read_rows(path: str):
    """Read the input file into a list of (line_number, fields, malformed)."""
    rows = []
    try:
        with open(path, encoding="utf-8") as f:
            for line_number, raw in enumerate(f, start=1):
                raw = raw.rstrip("\n").rstrip("\r")
                if not raw.strip():
                    continue
                fields, malformed = parse_line(raw)
                rows.append((line_number, fields, malformed))
    except UnicodeDecodeError as e:
        raise SystemExit(
            f"Could not read '{path}' as UTF-8 text: {e}\n"
            "The input must be the plain text file Agent 3 wrote. If this is "
            "already a .xlsx or another binary file, you have the arguments "
            "the wrong way round."
        )
    return rows


def build_workbook(rows: list):
    """Build the workbook. One sheet, header row, one row per input line."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Parts"

    header_font = Font(bold=True)
    for column, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=column, value=header)
        cell.font = header_font

    ws.freeze_panes = "A2"

    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    link_font = Font(color="0563C1", underline="single")

    for offset, (_line_number, fields, _malformed) in enumerate(rows):
        row = offset + 2
        part_info, price, url, image = fields

        ws.cell(row=row, column=1, value=part_info)

        price_value = parse_price(price)
        price_cell = ws.cell(row=row, column=2, value=price_value)
        if isinstance(price_value, float):
            # Display only. "5.00" parses to 5.0 and would otherwise show as
            # a bare "5"; the stored value is untouched either way.
            price_cell.number_format = "0.00"

        url_cell = ws.cell(row=row, column=3, value=url)
        if url.startswith("http://") or url.startswith("https://"):
            url_cell.hyperlink = url
            url_cell.font = link_font

        ws.cell(row=row, column=4, value=image)

    return wb


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit(
            "Usage:\n"
            "  python3 make_xlsx.py <input.txt> <output.xlsx>\n\n"
            "  <input.txt>   pipe-delimited text from Agent 3, one row per line:\n"
            "                  part info | price | url | image\n"
            "  <output.xlsx> spreadsheet to write (overwritten if it exists)"
        )

    input_path, output_path = sys.argv[1], sys.argv[2]

    if not os.path.exists(input_path):
        raise SystemExit(f"Input file not found: {input_path}")
    if not os.path.isfile(input_path):
        raise SystemExit(f"Input path is not a file: {input_path}")

    rows = read_rows(input_path)
    if not rows:
        raise SystemExit(
            f"Input file '{input_path}' contains no rows — nothing to convert."
        )

    malformed = [line_number for line_number, _fields, bad in rows if bad]
    for line_number in malformed:
        print(
            f"WARNING: line {line_number} does not have {FIELD_COUNT} fields. "
            f"Written to the sheet marked {MALFORMED_PREFIX}, not dropped.",
            file=sys.stderr,
        )

    wb = build_workbook(rows)
    try:
        wb.save(output_path)
    except OSError as e:
        raise SystemExit(f"Could not write '{output_path}': {e}")

    print(
        f"Wrote {output_path} — {len(rows)} row(s)"
        + (f", {len(malformed)} malformed" if malformed else "")
    )


if __name__ == "__main__":
    main()
